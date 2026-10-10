"""Aşama 12: proje raporunu (Word) üretir.

Rapordaki bütün sayılar outputs/ altındaki çıktılardan ve aynı hesap fonksiyonlarından okunur;
varsayımlar değişince betiği yeniden çalıştırmak raporu günceller. Word kuruluysa içindekiler
tablosu güncellenir ve raporun PDF kopyası da yazılır.

Çalıştırma:  py src/rapor_olustur.py   (önce hepsini_calistir.py çalışmış olmalı)
"""
import hashlib
import json
import subprocess
import sys
from datetime import datetime
from importlib.metadata import version

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import ayarlar as A
from optimizasyon import butce_ozeti, kampanya_listesi, kisitsiz_plan, optimize_et
from ortak import baslik_yaz, para, tr, yuzde
from senaryo_karsilastirma import ETKI_DUZEYLERI, duyarlilik_tablosu, senaryo_tablosu
from teklif_senaryosu import teklif_tablosu
from veri_hazirlama import ACIKLAMA, ham_veriyi_oku, kalite_kontrolu
from yonetici_ozeti import ozet_metni

KOYU_MAVI, IKINCIL, CIZGI, BASLIK_DOLGU = "0D366B", "52514E", "C3C2B7", "EEF3FA"
YAZILIMLAR = ["pandas", "numpy", "scikit-learn", "matplotlib", "PuLP", "highspy", "streamlit", "python-docx"]
CINSIYET = {"Female": "Kadın", "Male": "Erkek"}
YAS = {"0": "65 yaş altı", "1": "65 yaş ve üzeri"}


# --- Word yardımcıları -------------------------------------------------------------
class Rapor:
    """python-docx belgesini saran, başlık / paragraf / tablo / şekil ekleyen küçük yardımcı."""

    def __init__(self):
        self.doc = Document()
        self.sekil_no = self.tablo_no = 0
        self.no = {}  # metinde anılan tabloların numaraları
        bolum = self.doc.sections[0]
        bolum.page_width, bolum.page_height = Cm(21), Cm(29.7)  # A4
        bolum.left_margin = bolum.right_margin = bolum.top_margin = bolum.bottom_margin = Cm(2.5)

        normal = self.doc.styles["Normal"]
        normal.font.name, normal.font.size = "Calibri", Pt(11)
        normal.paragraph_format.space_after = Pt(7)
        normal.paragraph_format.line_spacing = 1.15
        dil = OxmlElement("w:lang")
        dil.set(qn("w:val"), "tr-TR")
        normal.element.get_or_add_rPr().append(dil)
        for ad, punto in (("Heading 1", 16), ("Heading 2", 12.5)):
            stil = self.doc.styles[ad]
            stil.font.name, stil.font.size, stil.font.bold = "Calibri", Pt(punto), True
            stil.font.color.rgb = RGBColor.from_string(KOYU_MAVI)
            for tema in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):  # şablonun tema yazı tipini bırak
                stil.element.rPr.rFonts.attrib.pop(qn(f"w:{tema}"), None)
        self.doc.styles["Heading 1"].paragraph_format.space_before = Pt(18)
        self.doc.styles["Heading 1"].paragraph_format.space_after = Pt(8)
        self.doc.styles["Heading 2"].paragraph_format.space_before = Pt(12)
        self.doc.styles["Heading 2"].paragraph_format.space_after = Pt(5)
        aciklama = self.doc.styles["Caption"]
        aciklama.font.size, aciklama.font.bold, aciklama.font.italic = Pt(9.5), False, True
        aciklama.font.color.rgb = RGBColor.from_string(IKINCIL)

        bolum.different_first_page_header_footer = True  # kapakta sayfa numarası yok
        alt_bilgi = bolum.footer.paragraphs[0]
        alt_bilgi.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._alan(alt_bilgi, "PAGE")

    @staticmethod
    def _alan(paragraf, kod, yer_tutucu="1"):
        """Paragrafa bir Word alanı (sayfa numarası, içindekiler) ekler."""
        for tur, etiket, metin in (("begin", None, None), (None, "w:instrText", kod), ("separate", None, None),
                                   (None, "w:t", yer_tutucu), ("end", None, None)):
            oge = OxmlElement(etiket or "w:fldChar")
            if tur:
                oge.set(qn("w:fldCharType"), tur)
            else:
                oge.set(qn("xml:space"), "preserve")
                oge.text = metin
            paragraf.add_run()._r.append(oge)

    def baslik(self, metin, duzey=1):
        self.doc.add_heading(metin, level=duzey)

    def p(self, metin, giris=None, italik=False, girinti=False):
        """Paragraf. `giris` verilirse paragrafın başına kalın yazılır."""
        paragraf = self.doc.add_paragraph()
        if giris:
            paragraf.add_run(giris + " ").bold = True
        paragraf.add_run(metin).italic = italik
        if girinti:
            paragraf.paragraph_format.left_indent = Cm(1)
        return paragraf

    def maddeler(self, *metinler):
        for metin in metinler:
            paragraf = self.doc.add_paragraph(metin, style="List Bullet")
            paragraf.paragraph_format.space_after = Pt(3)

    def sayfa_sonu(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def sekil(self, dosya, aciklama, genislik_cm=15):
        self.sekil_no += 1
        paragraf = self.doc.add_paragraph()
        paragraf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraf.paragraph_format.keep_with_next = True
        paragraf.paragraph_format.space_after = Pt(2)
        paragraf.add_run().add_picture(str(dosya), width=Cm(genislik_cm))
        alt = self.doc.add_paragraph(f"Şekil {self.sekil_no}. {aciklama}", style="Caption")
        alt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return self.sekil_no

    def tablo(self, aciklama, basliklar, satirlar, genislikler_cm, saga=(), ad=None):
        """Başlık satırı gölgeli tablo. `saga`: sağa hizalanacak sütunlar; `ad`: metinden anmak için etiket."""
        self.tablo_no += 1
        if ad:
            self.no[ad] = self.tablo_no
        ust = self.doc.add_paragraph(f"Tablo {self.tablo_no}. {aciklama}", style="Caption")
        ust.paragraph_format.keep_with_next = True
        ust.paragraph_format.space_after = Pt(3)

        tablo = self.doc.add_table(rows=1 + len(satirlar), cols=len(basliklar))
        tablo.style, tablo.alignment, tablo.autofit = "Table Grid", WD_TABLE_ALIGNMENT.CENTER, False
        kenarlar = OxmlElement("w:tblBorders")
        for kenar in ("top", "left", "bottom", "right", "insideH", "insideV"):
            oge = OxmlElement(f"w:{kenar}")
            for ozellik, deger in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", CIZGI)):
                oge.set(qn(f"w:{ozellik}"), deger)
            kenarlar.append(oge)
        tablo._tbl.tblPr.append(kenarlar)

        for i, satir in enumerate([basliklar] + [list(s) for s in satirlar]):
            for j, deger in enumerate(satir):
                hucre = tablo.cell(i, j)
                hucre.width = Cm(genislikler_cm[j])
                paragraf = hucre.paragraphs[0]
                paragraf.paragraph_format.space_after = Pt(0)
                paragraf.paragraph_format.line_spacing = 1.0
                if j in saga:
                    paragraf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                yazi = paragraf.add_run(str(deger))
                yazi.font.size, yazi.bold = Pt(9.5), i == 0
                if i == 0:
                    dolgu = OxmlElement("w:shd")
                    for ozellik, d in (("val", "clear"), ("color", "auto"), ("fill", BASLIK_DOLGU)):
                        dolgu.set(qn(f"w:{ozellik}"), d)
                    hucre._tc.get_or_add_tcPr().append(dolgu)
        for satir in tablo.rows:  # satır sayfa sonunda ikiye bölünmesin
            satir._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        tekrar = OxmlElement("w:tblHeader")  # başlık satırı her sayfada yinelensin
        tablo.rows[0]._tr.get_or_add_trPr().append(tekrar)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return self.tablo_no


def word_ile_tamamla(dosya):
    """Word kuruluysa içindekiler tablosunu günceller ve PDF kopyasını yazar; değilse sessizce geçer."""
    pdf = dosya.with_suffix(".pdf")
    # -20 ve -21: Word'ün yerleşik "İçindekiler 1" ve "İçindekiler 2" stilleri; satır aralığı sıkılaştırılır.
    betik = f"""
        $w = New-Object -ComObject Word.Application; $w.Visible = $false; $w.DisplayAlerts = 0
        try {{
            $d = $w.Documents.Open('{dosya}')
            foreach ($stil in -20, -21) {{
                $b = $d.Styles.Item($stil).ParagraphFormat; $b.SpaceBefore = 0; $b.SpaceAfter = 2; $b.LineSpacingRule = 0
            }}
            foreach ($t in $d.TablesOfContents) {{ $t.Update() }}
            $d.Save(); $d.ExportAsFixedFormat('{pdf}', 17); $d.Close($false)
        }} finally {{ if ($w.Documents.Count -eq 0) {{ $w.Quit() }} }}
    """
    try:
        sonuc = subprocess.run(["powershell", "-NoProfile", "-Command", betik], capture_output=True, timeout=180)
        return sonuc.returncode == 0 and pdf.exists()
    except (OSError, subprocess.TimeoutExpired):
        return False


# --- Veriler -----------------------------------------------------------------------
def verileri_topla():
    """Raporda kullanılan bütün tabloları ve sayıları tek sözlükte toplar."""
    ham = ham_veriyi_oku()
    temiz = pd.read_csv(A.TEMIZ_VERI)
    risk = pd.read_csv(A.RISK_TABLOSU)
    teklifler = teklif_tablosu(risk)
    liste = kampanya_listesi(risk, teklifler, optimize_et(teklifler))
    senaryolar = senaryo_tablosu(risk, teklifler, A.BUTCE)
    ayrildi = temiz[A.HEDEF].eq("Yes")
    return {
        "kontrol": kalite_kontrolu(ham),
        "temiz": temiz,
        "kesif": pd.read_csv(A.TABLO / "kesif_kayip_oranlari.csv").set_index(["degisken", "grup"]),
        "internet": ayrildi.groupby(temiz["InternetService"]).mean(),
        "metrik": json.loads(A.MODEL_METRIKLERI.read_text(encoding="utf-8")),
        "risk": risk,
        "gruplar": pd.read_csv(A.TABLO / "grup_bazli_kontrol.csv", dtype={"grup": str}),
        "teklifler": teklifler,
        "liste": liste,
        "ozet": butce_ozeti(liste, A.BUTCE),
        "kisitsiz": kisitsiz_plan(teklifler),
        "yillik_gelir": risk["MonthlyCharges"].sum() * 12,
        "senaryolar": senaryolar,
        "duyarlilik": duyarlilik_tablosu(senaryolar),
        "egri": pd.read_csv(A.TABLO / "butce_egrisi.csv", index_col="butce"),
    }


# --- Bölümler ----------------------------------------------------------------------
def kapak_ve_icindekiler(r):
    for _ in range(5):
        r.doc.add_paragraph()
    for metin, punto, kalin, renk in (
        ("KARAR DESTEK SİSTEMİ", 26, True, KOYU_MAVI),
        ("Müşteri Kaybı Riskini Tahmin Etme ve Kampanya Bütçesini Optimize Etme", 16, False, "0B0B0B"),
        ("Proje Raporu", 14, False, IKINCIL),
    ):
        paragraf = r.doc.add_paragraph()
        paragraf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        yazi = paragraf.add_run(metin)
        yazi.font.size, yazi.bold = Pt(punto), kalin
        yazi.font.color.rgb = RGBColor.from_string(renk)
    for _ in range(6):
        r.doc.add_paragraph()
    for metin in (f"Hazırlayan: {A.RAPOR_HAZIRLAYAN}", f"Ders: {A.RAPOR_DERS}",
                  f"Öğretim elemanı: {A.RAPOR_OGRETIM_ELEMANI}", datetime.now().strftime("%d.%m.%Y")):
        paragraf = r.doc.add_paragraph(metin)
        paragraf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraf.paragraph_format.space_after = Pt(2)
    r.sayfa_sonu()

    paragraf = r.doc.add_paragraph()
    yazi = paragraf.add_run("İçindekiler")
    yazi.font.size, yazi.bold = Pt(16), True
    yazi.font.color.rgb = RGBColor.from_string(KOYU_MAVI)
    r._alan(r.doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u',
            "İçindekiler tablosu için bu satıra sağ tıklayıp “Alanı Güncelleştir” seçeneğini kullanın.")
    r.sayfa_sonu()


def yonetici_ozeti(r, v):
    r.baslik("1. Yönetici özeti")
    r.p("Bu proje, abonelik modeliyle çalışan bir telekom işletmesinin iki sorusuna birlikte yanıt veren bir karar destek "
        "sistemi prototipidir: hangi müşterilerin ayrılma riski yüksektir ve sınırlı kampanya bütçesi hangi müşteriye, "
        "hangi teklifle harcanmalıdır?")
    r.p("Sistem dört katmandan oluşur: müşteri kaybı tahmini, her müşteri için teklif seçeneklerinin ekonomik "
        "değerlendirmesi, bütçe kısıtı altında optimizasyon ve yöneticiye öneriyi gerekçesi ve belirsizliğiyle birlikte "
        "sunan karar ekranı.")
    for paragraf in ozet_metni(v["risk"], v["ozet"], v["senaryolar"], v["metrik"]["secilen_esik"]).split("\n\n"):
        r.p(paragraf)
    r.p("Müşteri verisi gerçektir (IBM Telco Customer Churn). Teklif maliyetleri, tekliflerin etkisi, bütçe ve kapasite "
        "ise sentetik varsayımdır; veri setinde kampanya bilgisi yoktur. Bu nedenle net fayda sayıları yöntemin nasıl "
        "çalıştığını gösterir, gerçek bir kampanyanın sonucunu kanıtlamaz.", giris="Temel sınır:")


def problem(r):
    r.baslik("2. Problem tanımı ve kapsam")
    r.p("Abonelik tabanlı işletmeler müşterilerinin bir bölümünü fiyat, rakip teklifleri, hizmet kalitesi veya müşteri "
        "hizmetleri deneyimi nedeniyle kaybeder. İşletme iptal başvurusundan sonra harekete geçerse müşteriyi geri kazanmak "
        "zorlaşır. Öte yandan herkese indirim vermek de doğru değildir: bazı müşteriler zaten kalacaktır, bazıları indirim "
        "verilse de ayrılacaktır. Bütçenin yanlış müşterilere harcanması kârlılığı düşürür.")
    r.p("Karar vericinin yanıt aradığı sorular şunlardır:")
    r.maddeler(
        "Hangi müşteriler için ayrılma riski yüksektir?",
        "Riskli müşteride öne çıkan işaretler nelerdir: sözleşme türü mü, ücret mi, hizmet mi?",
        "Hangi müdahale seçeneği bu müşteri için ekonomik olarak anlamlıdır?",
        "Kampanya bütçesi hangi müşteri ve tekliflere ayrılmalıdır?",
        "Önerinin beklenen faydası nedir ve ne kadar belirsizlik içerir?",
    )
    r.baslik("2.1. Kapsam ve sınırlar", 2)
    r.p("Prototip telekomünikasyon alanı için kurulmuştur. Müşteri kaybı tahmini açık bir eğitim veri setiyle yapılır. "
        "Kampanya etkisi ve teklif maliyetleri için gerçek veri bulunmadığından ayrı ve açıkça “sentetik” olarak "
        "işaretlenmiş bir senaryo tablosu kullanılır. Belirli bir teklifin müşteriyi kesin olarak elde tutacağı iddia "
        "edilmez; sistem nihai kararı vermez, yöneticiye öneri sunar.")
    r.baslik("2.2. Başarı ölçütleri", 2)
    r.maddeler(
        "Tahmin katmanı: görülmemiş test verisinde duyarlılık, kesinlik, F1, ROC-AUC, PR-AUC ve kalibrasyon.",
        "Karar katmanı: aynı bütçe ve aynı varsayımlar altında beklenen net faydanın mevcut yöntemle karşılaştırılması.",
        "Şeffaflık: sentetik varsayımların gerçek verilerden ayrılması ve sonuçların varsayımlara duyarlılığının gösterilmesi.",
    )


def veri(r, v):
    k, dosya = v["kontrol"], A.HAM_VERI
    r.baslik("3. Veri kaynağı ve veri sözlüğü")
    r.baslik("3.1. Kaynak kaydı", 2)
    r.p("Ana veri seti IBM Telco Customer Churn örnek verisidir. Her satır bir müşteriyi temsil eder; hizmetler, abonelik "
        "süresi, sözleşme ve faturalandırma bilgilerini ve müşterinin ayrılıp ayrılmadığını gösteren Churn sütununu içerir.")
    r.tablo("Veri kaynağı kaydı", ["Bilgi", "Değer"], [
        ("Veri seti", "IBM Telco Customer Churn"),
        ("İndirilen kaynak", A.INDIRILEN_KAYNAK),
        ("Dosya", f"data/raw/{dosya.name}"),
        ("Dosya tarihi (indirme)", datetime.fromtimestamp(dosya.stat().st_mtime).strftime("%d.%m.%Y %H:%M")),
        ("Boyut", f"{tr(dosya.stat().st_size)} bayt; {tr(k['satir'])} satır × {k['sutun']} sütun"),
        ("SHA-256", hashlib.sha256(dosya.read_bytes()).hexdigest()),
    ], [4.2, 11.8], ad="kaynak")
    r.p("Veri seti geçmiş müşteri özellikleri ile ayrılma durumu arasındaki örüntüleri incelemeye uygundur. Hangi müşteriye "
        "hangi kampanyanın sunulduğunu ve kampanyanın sonucu değiştirip değiştirmediğini göstermez. Bu nedenle yalnızca "
        "müşteri kaybı tahmininde kullanılmıştır; kampanya etkisini kanıtlamaz. Veri kesitsel bir örnektir ve gerçek bir "
        "işletmenin güncel müşterilerini temsil etmeyebilir.", giris="Veri seti neyi gösterir, neyi göstermez?")
    r.baslik("3.2. Veri sözlüğü", 2)
    satirlar = []
    for sutun, aciklama in ACIKLAMA.items():
        if sutun == A.HEDEF:
            rol = "Hedef değişken; girdi değildir."
        elif sutun in A.MODEL_DISI:
            rol = "Girdi değil. " + A.MODEL_DISI[sutun]
        else:
            rol = "Girdi."
        satirlar.append((sutun, aciklama, rol))
    r.tablo("Veri sözlüğü ve sütunların modeldeki rolü", ["Sütun", "Açıklama", "Modeldeki rolü"], satirlar, [3.3, 5.7, 7.0], ad="sozluk")
    r.p("Demografik alanlar (cinsiyet, yaş grubu, eş ve bakmakla yükümlü olunan kişi bilgisi) ayrımcılık riski taşıdığı ve "
        "iş gerekçesi zayıf olduğu için model girdisi yapılmamıştır. Cinsiyet ve yaş grubu yalnızca sonuçların gruplar "
        "arasında nasıl dağıldığını denetlemek için kullanılmıştır (Bölüm 6.7).")


def kalite(r, v):
    k, temiz = v["kontrol"], v["temiz"]
    churn = k["churn_dagilimi"]
    r.baslik("4. Veri kalite kontrolü ve temizleme")
    r.p("İndirilen dosya modele doğrudan verilmemiş, önce aşağıdaki kontroller yapılmıştır. Kontroller ve dönüşümler "
        "src/veri_hazirlama.py betiğiyle tekrarlanabilir; ham dosya değiştirilmez.")
    r.maddeler(
        f"Boyut: {tr(k['satir'])} satır ve {k['sutun']} sütun; kaynak açıklamasındaki boyutla "
        f"({tr(A.BEKLENEN_SATIR)} × {A.BEKLENEN_SUTUN}) uyumludur.",
        f"Tekrar: customerID alanında tekrar eden kayıt sayısı {k['tekrar_eden_kimlik']}, tümüyle aynı satır sayısı {k['tekrar_eden_satir']}.",
        f"Eksik değer: TotalCharges ham dosyada metin olarak gelir ve {k['bos_toplam_ucret']} kayıtta boştur. Bu kayıtların "
        f"tümünde tenure değeri {k['bos_toplam_ucret_tenure']} olduğundan (henüz faturası oluşmamış yeni müşteriler) boş değerler "
        "0 ile doldurulmuş ve ayrı bir sütunla işaretlenmiştir. Başka boş hücre yoktur; satır silinmemiştir.",
        f"Sınıf dağılımı: {tr(churn['Yes'])} müşteri ayrılmış, {tr(churn['No'])} müşteri kalmıştır; ayrılma oranı "
        f"{yuzde(churn['Yes'] / k['satir'])}. Sınıflar dengesiz olduğu için değerlendirmede yalnızca doğruluk oranına bakılmamıştır.",
        "Tip dönüşümü: tenure ve SeniorCitizen tam sayıya, MonthlyCharges ve TotalCharges ondalık sayıya çevrilmiştir.",
    )
    r.tablo(f"Sayısal alanların özeti (temiz veri; ücretler {A.PARA})", ["Alan", "En küçük", "Ortanca", "Ortalama", "En büyük"],
            [(s, tr(temiz[s].min(), 2), tr(temiz[s].median(), 2), tr(temiz[s].mean(), 2), tr(temiz[s].max(), 2))
             for s in A.SAYISAL], [4.0, 3.0, 3.0, 3.0, 3.0], saga=(1, 2, 3, 4))


def kesif(r, v):
    k = v["kesif"]["kayip_orani"]
    genel = v["kontrol"]["churn_dagilimi"]["Yes"] / v["kontrol"]["satir"]
    r.baslik("5. Keşifsel analiz")
    r.p("Model kurulmadan önce kayıp oranı başlıca müşteri gruplarında karşılaştırılmıştır. Grafiklerdeki yatay ya da dikey "
        f"gri çizgi genel kayıp oranını ({yuzde(genel)}) gösterir.")
    r.sekil(A.GRAFIK / "01_genel_kayip_orani.png", "Genel müşteri kayıp oranı", 13)
    r.sekil(A.GRAFIK / "02_sozlesme_turune_gore.png", "Sözleşme türüne göre kayıp oranı")
    r.p(f"En belirgin fark sözleşme türündedir: aylık (taahhütsüz) sözleşmede kayıp oranı {yuzde(k['Sözleşme türü', 'Month-to-month'])}, "
        f"bir yıllıkta {yuzde(k['Sözleşme türü', 'One year'])}, iki yıllıkta {yuzde(k['Sözleşme türü', 'Two year'])} düzeyindedir.")
    r.sekil(A.GRAFIK / "03_abonelik_suresine_gore.png", "Abonelik süresine göre kayıp oranı")
    r.p(f"Kayıp ilk yılda yoğunlaşır: ilk 12 ayındaki müşterilerde oran {yuzde(k['Abonelik süresi', '0–12 ay'])}, "
        f"49–72 aydır müşteri olanlarda {yuzde(k['Abonelik süresi', '49–72 ay'])} düzeyindedir.")
    r.sekil(A.GRAFIK / "04_aylik_ucrete_gore.png", "Aylık ücret aralıklarına göre kayıp oranı")
    r.p(f"Aylık ücreti 35 {A.PARA} altında olan müşterilerde kayıp oranı {yuzde(k['Aylık ücret', A.UCRET_ETIKET[0]])} ile en düşüktür; "
        f"75–95 {A.PARA} aralığında {yuzde(k['Aylık ücret', '75–95'])} ile en yüksek düzeye çıkar. İlişki tek yönlü artan değildir.")
    r.sekil(A.GRAFIK / "05_teknik_destege_gore.png", "Teknik destek durumuna göre kayıp oranı")
    r.p(f"Teknik destek hizmeti almayan internet müşterilerinde kayıp oranı {yuzde(k['Teknik destek', 'No'])}, alanlarda "
        f"{yuzde(k['Teknik destek', 'Yes'])} düzeyindedir.")
    r.p("Bu grafikler ilişki gösterir, neden-sonuç ilişkisini kanıtlamaz. Örneğin teknik destek almayanlarda kaybın yüksek "
        "olması, desteğin kendisinden değil bu müşterilerin başka özelliklerinden (sözleşme türü, abonelik süresi) "
        "kaynaklanıyor olabilir. Bu ayrım, teklif etkilerinin neden veriden türetilmeyip varsayım olarak bırakıldığını da açıklar.",
        giris="Yorum sınırı:")


def model(r, v):
    m, risk = v["metrik"], v["risk"]
    esik = m["secilen_esik"]
    loj, agac = "Lojistik regresyon", "Karar ağacı"
    r.baslik("6. Tahmin modeli")
    r.baslik("6.1. Girdiler", 2)
    r.p(f"Hedef değişken Churn sütunudur. Modelde {len(A.SAYISAL)} sayısal ({', '.join(A.SAYISAL)}) ve {len(A.KATEGORIK)} kategorik "
        "alan (hizmetler, sözleşme türü, faturalandırma ve ödeme yöntemi) girdi olarak kullanılmıştır. customerID ve dört "
        f"demografik alan girdi değildir (Tablo {r.no['sozluk']}).")
    r.baslik("6.2. Yöntem", 2)
    r.maddeler(
        f"Veri katmanlı (stratified) biçimde %{round((1 - A.TEST_ORANI) * 100)} eğitim ve %{round(A.TEST_ORANI * 100)} test olarak bölünmüştür: "
        f"{tr(m['veri']['egitim'])} eğitim, {tr(m['veri']['test'])} test müşterisi. Rastgelelik tohumu {m['rastgele_tohum']} olarak kaydedilmiştir.",
        "Ölçekleme ve kategorik kodlama model ile aynı işlem hattındadır; dönüşümler yalnızca eğitim verisinden öğrenilir ve test verisine aynen uygulanır.",
        "Lojistik regresyon temel model, karar ağacı karşılaştırma modelidir. Az sayıda modeli doğru değerlendirmek tercih edilmiştir.",
        f"Model ayarları eğitim verisi içinde {A.CV_KAT} katlı çapraz doğrulamayla (ölçüt: ROC-AUC) seçilmiştir. Seçilen ayarlar: "
        f"lojistik regresyon C = {tr(m['en_iyi_parametreler'][loj]['model__C'], 2)}; karar ağacı en büyük derinlik "
        f"{m['en_iyi_parametreler'][agac]['model__max_depth']}, yaprak başına en az {m['en_iyi_parametreler'][agac]['model__min_samples_leaf']} müşteri.",
        "Test verisi model ve eşik seçiminde kullanılmamış, yalnızca en sonda bir kez değerlendirme için kullanılmıştır.",
    )
    r.baslik("6.3. Sonuçlar", 2)
    satirlar = [("ROC-AUC (çapraz doğrulama)",) + tuple(tr(m["capraz_dogrulama"][ad]["roc_auc"], 3) for ad in (loj, agac)),
                ("PR-AUC (çapraz doğrulama)",) + tuple(tr(m["capraz_dogrulama"][ad]["pr_auc"], 3) for ad in (loj, agac))]
    for etiket, anahtar in (("ROC-AUC (test)", "roc_auc"), ("PR-AUC (test)", "pr_auc"), ("Brier skoru (test)", "brier")):
        satirlar.append((etiket,) + tuple(tr(m["test"][ad][anahtar], 3) for ad in (loj, agac)))
    for blok, aciklama in (("secilen_esik", "seçilen eşik"), ("esik_0_50", "eşik 0,50")):
        satirlar.append((f"Eşik ({aciklama})",) + tuple(tr(m["test"][ad][blok]["esik"], 2) for ad in (loj, agac)))
        for etiket, anahtar in (("Kesinlik (precision)", "kesinlik"), ("Duyarlılık (recall)", "duyarlilik"), ("Doğruluk", "dogruluk")):
            satirlar.append((f"{etiket}, {aciklama}",) + tuple(yuzde(m["test"][ad][blok][anahtar]) for ad in (loj, agac)))
        satirlar.append((f"F1, {aciklama}",) + tuple(tr(m["test"][ad][blok]["f1"], 3) for ad in (loj, agac)))
    r.tablo("Model değerlendirme ölçütleri", ["Ölçüt", loj, agac], satirlar, [7.6, 4.2, 4.2], saga=(1, 2), ad="olcut")
    s = m["test"][loj]["secilen_esik"]
    (dn, yp), (yn, dp) = s["karisiklik_matrisi"]
    r.p(f"Lojistik regresyon test verisinde ROC-AUC {tr(m['test'][loj]['roc_auc'], 3)} ile karar ağacından "
        f"({tr(m['test'][agac]['roc_auc'], 3)}) biraz daha iyi sıralama yapmaktadır; çapraz doğrulama sonuçları da aynı yöndedir. "
        f"Seçilen eşikte ayrılan müşterilerin {yuzde(s['duyarlilik'])} kadarı yakalanmakta, riskli olarak işaretlenenlerin "
        f"{yuzde(s['kesinlik'])} kadarı gerçekten ayrılmaktadır.")
    r.p(f"Doğruluk oranı tek başına yanıltıcıdır: hiçbir müşteriyi riskli saymayan bir kural test verisinde "
        f"{yuzde(1 - m['veri']['test_kayip_orani'])} doğruluk elde eder ama ayrılan hiçbir müşteriyi yakalamaz. Bu nedenle ölçütler birlikte raporlanmıştır.")
    r.tablo(f"Lojistik regresyonun karışıklık matrisi (test verisi, eşik {tr(s['esik'], 2)})",
            ["", "Tahmin: kalır", "Tahmin: ayrılır"],
            [("Gerçek: kaldı", tr(dn), tr(yp)), ("Gerçek: ayrıldı", tr(yn), tr(dp))], [5.0, 4.0, 4.0], saga=(1, 2))
    r.sekil(A.GRAFIK / "06_roc_pr_egrileri.png", "Modellerin ROC ve kesinlik–duyarlılık eğrileri (test verisi)", 16)
    r.baslik("6.4. Kalibrasyon", 2)
    r.p("Ekonomik hesaplarda modelin ürettiği olasılık doğrudan kullanıldığı için olasılıkların gerçek oranlarla uyumu "
        "önemlidir. Test müşterileri tahmin edilen olasılığa göre on eşit gruba ayrılmış ve her grupta gerçekleşen kayıp "
        f"oranı tahminle karşılaştırılmıştır. Lojistik regresyonun Brier skoru {tr(m['test'][loj]['brier'], 3)} düzeyindedir; "
        "noktalar tam uyum çizgisine yakındır, orta olasılık aralığında model riski bir miktar yüksek tahmin etmektedir.")
    r.sekil(A.GRAFIK / "07_kalibrasyon.png", "Tahmin edilen olasılıkların kalibrasyonu (test verisi)", 11)
    r.baslik("6.5. Eşik seçimi ve risk grupları", 2)
    r.p(f"Sınıflandırma eşiği {tr(esik, 2)} olarak seçilmiştir. Bu değer, eğitim verisindeki çapraz doğrulama tahminlerinde F1 "
        "skorunu en yüksek yapan olasılıktır. Ayrılan bir müşteriyi kaçırmak gereksiz bir teklif sunmaktan genellikle daha "
        "maliyetli olduğu için 0,50 yerine daha düşük bir eşik tercih edilmiş, böylece duyarlılık artmıştır. Risk grupları bu eşiğe göre tanımlanmıştır.")
    gruplar = risk.groupby("risk_grubu")["gercek_churn"].agg(["size", "mean"])
    araliklar = {"Yüksek": f"{tr(esik, 2)} ve üzeri", "Orta": f"{tr(esik / 2, 2)} – {tr(esik, 2)}", "Düşük": f"{tr(esik / 2, 2)} altı"}
    r.tablo("Risk grupları ve gruplarda gerçekleşen kayıp oranı (test verisi)",
            ["Risk grubu", "Ayrılma olasılığı", "Müşteri", "Gerçekleşen kayıp oranı"],
            [(g, araliklar[g], tr(gruplar.loc[g, "size"]), yuzde(gruplar.loc[g, "mean"])) for g in ("Yüksek", "Orta", "Düşük")],
            [3.4, 4.6, 3.0, 5.0], saga=(2, 3))
    r.p("Eşik yalnızca raporlama ve risk grupları içindir. Kampanya kararında eşik kullanılmaz; her müşterinin olasılığı "
        "doğrudan ekonomik hesaba girer (Bölüm 7).")
    r.baslik("6.6. Risk işaretleri", 2)
    r.p("Yöneticiye her müşteri için en fazla üç risk işareti gösterilir. İşaretler, müşterinin içinde bulunduğu ve eğitim "
        "verisinde kayıp oranı genel oranın en az 1,25 katı olan gruplardır (örneğin aylık sözleşme, ilk 12 ay, elektronik "
        "çekle ödeme). Bu betimleyici bir özettir; modelin iç hesabının dökümü değildir ve ayrılmanın nedenini kanıtlamaz.")
    r.p("Lojistik regresyon katsayıları bu amaçla kullanılmamıştır: aylık ücret, hizmet sütunlarının neredeyse doğrusal bir "
        "bileşimi olduğundan katsayılar tek tek yorumlanamamakta, ilk denemede yanıltıcı gerekçeler üretmiştir.")
    r.baslik("6.7. Grup bazlı kontrol", 2)
    g = v["gruplar"]
    adlar = {"gender": ("Cinsiyet", CINSIYET), "SeniorCitizen": ("Yaş grubu", YAS)}
    r.tablo("Model girdisi olmayan demografik gruplarda sonuçlar (test verisi)",
            ["Değişken", "Grup", "Müşteri", "Gerçek kayıp", "Ortalama tahmin", "Yüksek risk payı", "Duyarlılık", "Kesinlik"],
            [(adlar[s["grup_degiskeni"]][0], adlar[s["grup_degiskeni"]][1][s["grup"]], tr(s["musteri_sayisi"]),
              yuzde(s["gercek_kayip_orani"]), yuzde(s["ortalama_tahmin"]), yuzde(s["yuksek_risk_payi"]),
              yuzde(s["duyarlilik"]), yuzde(s["kesinlik"])) for _, s in g.iterrows()],
            [2.0, 2.9, 1.6, 1.9, 2.0, 2.0, 1.9, 1.7], saga=(2, 3, 4, 5, 6, 7))
    yasli, genc = g[g["grup"] == "1"].iloc[0], g[g["grup"] == "0"].iloc[0]
    r.p("Cinsiyete göre sonuçlar birbirine yakındır. Yaş grubu model girdisi olmadığı hâlde 65 yaş ve üzeri müşterilerin "
        f"{yuzde(yasli['yuksek_risk_payi'])} kadarı yüksek riskli grupta yer almaktadır; bu oran diğer müşterilerde "
        f"{yuzde(genc['yuksek_risk_payi'])} düzeyindedir. Fark, bu grubun gerçek kayıp oranının da yüksek olmasıyla "
        f"({yuzde(yasli['gercek_kayip_orani'])}) uyumludur ve sözleşme türü gibi girdiler üzerinden dolaylı olarak oluşur. "
        "Sistem gerçek müşterilerle kullanılacaksa bu grubun tekliflerden nasıl etkilendiği ayrıca izlenmelidir.")


def teklif(r, v):
    k, ic, kisitsiz = v["kesif"]["kayip_orani"], v["internet"], v["kisitsiz"]
    r.baslik("7. Teklif senaryosu (sentetik)")
    r.baslik("7.1. Neden ikinci bir tablo?", 2)
    r.p("Ayrılma olasılığı tek başına kampanya kararı değildir. Karar için her teklifin maliyeti ve müşteri davranışına ek "
        "etkisi gerekir; veri setinde bunların hiçbiri yoktur. Bu nedenle her müşteri ve teklif için maliyet ile etki "
        "varsayımlarını içeren ikinci bir tablo üretilmiş ve dosya adında açıkça işaretlenmiştir "
        f"({A.TEKLIF_TABLOSU.relative_to(A.KOK).as_posix()}).")
    r.baslik("7.2. Ekonomik hesap", 2)
    r.p("Dört seçenek karşılaştırılır: indirim, uygun paket, teknik destek ve teklif sunmama. Her müşteri–teklif çifti için:")
    for satir in ("Ek elde tutma olasılığı = ayrılma olasılığı × teklifin etki oranı",
                  f"Elde tutulan müşterinin katkısı = aylık ücret × katkı marjı × {A.UFUK_AY} ay",
                  "Beklenen ek katkı = ek elde tutma olasılığı × elde tutulan müşterinin katkısı",
                  "Beklenen net fayda = beklenen ek katkı − teklif maliyeti"):
        r.p(satir, italik=True, girinti=True).paragraph_format.space_after = Pt(2)
    r.p("Etki oranı, “teklif sunulmasaydı ayrılacak olan müşterilerin ne kadarı teklifle kalır” sorusunun varsayılan "
        "yanıtıdır. Böylece müşterinin tüm geliri kampanyanın başarısı sayılmaz: zaten kalacak bir müşteriye verilen teklif "
        "yalnızca maliyet yaratır. Teklif maliyetinin, müşteri sonradan ayrılsa bile tamamının harcandığı kabul edilmiştir "
        "(ihtiyatlı varsayım).")
    r.baslik("7.3. Varsayımlar ve dayanakları", 2)
    r.p("Aşağıdaki değerlerin hiçbiri gerçek bir işletmeden, sektör kaynağından veya kampanya deneyinden alınmamıştır. "
        "Tekliflerin yapısı proje yol haritasından, kime uygun oldukları ve kimde daha etkili sayıldıkları keşifsel "
        "analizdeki örüntülerden gelir. Sayıların kendisi yuvarlak ve makul görünen senaryo parametreleridir.")
    etki = lambda sozluk, anahtar: yuzde(sozluk[anahtar], 0)
    r.tablo("Senaryo varsayımları ve dayanakları", ["Varsayım", "Değer", "Dayanak"], [
        ("Para birimi", A.PARA, "Veri setindeki ücretlerin birimi; kur varsayımı eklememek için çevrilmemiştir."),
        ("Değerlendirme dönemi", f"{A.UFUK_AY} ay", "Yuvarlak varsayım: elde tutulan müşterinin bir yıllık katkısı."),
        ("Katkı marjı", f"Aylık ücretin {yuzde(A.KATKI_MARJI, 0)} kadarı",
         "Yuvarlak varsayım. Veri setinde maliyet bilgisi olmadığı için gelir yerine katkı kârı bu oranla temsil edilmiştir."),
        ("Kampanya bütçesi", para(A.BUTCE),
         f"Sınırsız bütçeyle en iyi planın maliyeti {para(kisitsiz['maliyet'])}; bütçe bunun yaklaşık yarısıdır ve portföyün yıllık "
         f"gelirinin {yuzde(A.BUTCE / v['yillik_gelir'])} kadarıdır. Bu düzeyde bütçe gerçekten kısıtlar."),
        ("Teknik destek kapasitesi", f"{A.DESTEK_KAPASITESI} müşteri",
         f"Sınır olmasa {tr(kisitsiz['teklif_dagilimi'][A.DESTEK])} müşteriye destek önerilir; kapasite bunun yaklaşık yarısıdır."),
        ("İndirim: maliyet", f"{A.INDIRIM_AY} ay boyunca {yuzde(A.INDIRIM_ORANI, 0)} indirim", "Yuvarlak varsayım."),
        ("İndirim: etki oranı", f"Aylık sözleşmede {etki(A.INDIRIM_ETKI, 'Month-to-month')}, taahhütlü sözleşmede {etki(A.INDIRIM_ETKI, 'One year')}",
         f"Yön keşifsel analizden: aylık sözleşmede kayıp {yuzde(k['Sözleşme türü', 'Month-to-month'])}. Sayılar varsayımdır."),
        ("Uygun paket: uygunluk", f"Aylık ücreti {A.PAKET_ALT_SINIR} {A.PARA} ve üzeri olanlar",
         "Düşük ücretli müşteri için daha ucuz bir paket olmadığı varsayılmıştır; yüksek ücret aralıklarında kayıp oranı daha yüksektir."),
        ("Uygun paket: maliyet", f"{A.PAKET_GECIS_MALIYETI} {A.PARA} geçiş + {A.UFUK_AY} ay {yuzde(A.PAKET_UCRET_DUSUSU, 0)} ücret düşüşü", "Yuvarlak varsayım."),
        ("Uygun paket: etki oranı", f"{A.PAKET_YUKSEK_UCRET} {A.PARA} ve üzerinde {yuzde(A.PAKET_ETKI_YUKSEK, 0)}, altında {yuzde(A.PAKET_ETKI_NORMAL, 0)}",
         "Varsayım. İndirim değerleriyle birlikte, her teklifin en az bir müşteri profilinde öne çıkacağı biçimde ayarlanmıştır."),
        ("Teknik destek: uygunluk", "İnterneti olup teknik destek almayanlar",
         f"Keşifsel analiz: teknik desteği olmayanlarda kayıp {yuzde(k['Teknik destek', 'No'])}, olanlarda {yuzde(k['Teknik destek', 'Yes'])} (ilişki, nedensellik değil)."),
        ("Teknik destek: maliyet", f"Müşteri başına {A.DESTEK_MALIYETI} {A.PARA}", "Yuvarlak varsayım."),
        ("Teknik destek: etki oranı", f"Fiberde {etki(A.DESTEK_ETKI, 'Fiber optic')}, DSL'de {etki(A.DESTEK_ETKI, 'DSL')}",
         f"Yön veriden: fiber kullananlarda kayıp {yuzde(ic['Fiber optic'])}, DSL kullananlarda {yuzde(ic['DSL'])}. Sayılar varsayımdır."),
        ("Kanıt düzeyi", "Tüm teklifler için düşük", "Hiçbir değer gerçek kampanya verisine dayanmamaktadır."),
    ], [3.6, 4.4, 8.0])
    r.p("İndirim ve uygun paket değerleri, her teklifin en az bir müşteri profilinde diğerlerinden iyi çıkacağı biçimde "
        "seçilmiştir; aksi hâlde tek bir teklif her müşteride kazanmakta ve teklif karşılaştırması anlamsızlaşmaktadır. "
        "Bu, sonuçların bu varsayımlara bağlı olduğunu açıkça gösterir. Bütün değerler src/ayarlar.py dosyasında tek yerde "
        "durur ve değiştirilip sistem yeniden çalıştırılabilir.", giris="Açık not:")
    r.baslik("7.4. Tablonun özeti", 2)
    gercek = v["teklifler"][v["teklifler"]["offer_type"] != A.TEKLIF_YOK]
    ozet = gercek.groupby("offer_type").agg(uygun=("customerID", "size"), maliyet=("offer_cost", "mean"),
                                            pozitif=("beklenen_net_fayda", lambda s: (s > 0).sum())).reindex(A.TEKLIFLER)
    r.tablo(f"Teklif türlerine göre uygunluk ve beklenen fayda ({tr(len(v['risk']))} müşteri)",
            ["Teklif", "Uygun olduğu müşteri", "Ortalama maliyet", "Beklenen net faydası pozitif olan müşteri"],
            [(t, tr(s["uygun"]), para(s["maliyet"], 2), tr(s["pozitif"])) for t, s in ozet.iterrows()],
            [3.4, 3.8, 3.4, 5.4], saga=(1, 2, 3))
    r.p("Müşteriye uygun olmayan teklif tabloda yer almaz ve önerilemez (örneğin interneti olmayan müşteriye teknik destek). "
        "Tablodaki alanlar: customerID, offer_type, offer_cost, estimated_incremental_retention, retained_customer_contribution, "
        "beklenen net fayda ve confidence_or_evidence_level.")


def optimizasyon(r, v):
    ozet, kisitsiz, liste = v["ozet"], v["kisitsiz"], v["liste"]
    onerilen = liste[liste["offer_type"] != A.TEKLIF_YOK]
    r.baslik("8. Bütçe optimizasyonu")
    r.baslik("8.1. Model", 2)
    r.p("Müşteri–teklif seçimi ikili karar değişkenli doğrusal bir optimizasyon modeliyle yapılır. Her müşteri–teklif çifti "
        "için bir karar değişkeni tanımlanır (1: teklif sunulur, 0: sunulmaz). Amaç toplam beklenen net faydayı en büyük "
        "yapmaktır. Kısıtlar:")
    r.maddeler(
        "Toplam teklif maliyeti kampanya bütçesini aşamaz.",
        "Bir müşteriye aynı dönemde en fazla bir teklif sunulur.",
        "Teknik destek sunulan müşteri sayısı ekip kapasitesini aşamaz.",
        "Müşteriye uygun olmayan teklif seçilemez.",
    )
    r.p("Beklenen net faydası pozitif olmayan seçenekler modele alınmaz: müşteri riskli diye zarar eden bir teklif "
        "sunulmaz. Model PuLP kütüphanesiyle kurulmuş, HiGHS çözücüsüyle en iyi çözüme kadar çözülmüştür.")
    r.baslik("8.2. Bütçe düzeyinin dayanağı", 2)
    r.p(f"Bütçe ve kapasite sınırı olmasaydı en iyi plan {tr(kisitsiz['teklif_sunulan_musteri'])} müşteriye teklif sunar, "
        f"{para(kisitsiz['maliyet'])} harcar ve {para(kisitsiz['beklenen_net_fayda'])} beklenen net fayda üretirdi. "
        f"{para(A.BUTCE)} bütçe bu tutarın yaklaşık yarısıdır. Daha düşük bütçelerde plan neredeyse yalnızca teknik destekten "
        "oluşmakta, çok daha yüksek bütçelerde ise faydalı tekliflerin hemen hepsi karşılandığı için seçim yapmaya gerek "
        "kalmamaktadır. Seçilen düzeyde bütçe gerçekten kısıtlar ve üç teklif türü de plana girer.")
    r.baslik("8.3. Önerilen plan", 2)
    r.p(f"Sistem {para(ozet['butce'])} bütçenin {para(ozet['harcama'])} tutarındaki bölümünü ({yuzde(ozet['kullanim_orani'])}) "
        f"{tr(ozet['teklif_sunulan_musteri'])} müşteriye ayırmayı önermektedir. Planın beklenen ek katkısı "
        f"{para(ozet['beklenen_ek_katki'])}, beklenen net faydası {para(ozet['beklenen_net_fayda'])} düzeyindedir.")
    dagilim = onerilen.groupby("offer_type").agg(n=("customerID", "size"), maliyet=("offer_cost", "sum"),
                                                net=("beklenen_net_fayda", "sum")).reindex(A.TEKLIFLER, fill_value=0)
    r.tablo("Önerilen planda teklif türlerinin dağılımı", ["Teklif", "Müşteri", "Toplam maliyet", "Beklenen net fayda"],
            [(t, tr(s["n"]), para(s["maliyet"]), para(s["net"])) for t, s in dagilim.iterrows()]
            + [("Toplam", tr(dagilim["n"].sum()), para(dagilim["maliyet"].sum()), para(dagilim["net"].sum()))],
            [4.0, 3.0, 4.0, 4.5], saga=(1, 2, 3))
    r.maddeler(
        f"Önerilen {tr(ozet['varsayima_duyarli_teklif'])} teklif, gerçek etki varsayılanın yarısı kadar olursa maliyetini "
        "karşılamamaktadır; bu teklifler listede “Varsayıma duyarlı” olarak işaretlenmiştir.",
        f"Beklenen net faydası pozitif olan {tr(ozet['butce_disinda_kalan'])} müşteri, bütçe veya teknik destek kapasitesi "
        "yetmediği için liste dışında kalmıştır.",
        f"Kalan {tr(len(liste) - len(onerilen) - ozet['butce_disinda_kalan'])} müşteride hiçbir teklifin beklenen net faydası pozitif değildir; "
        "bu müşterilere teklif sunulmaması önerilir.",
    )
    r.tablo("Önerilen kampanya listesinden ilk on müşteri (beklenen net faydaya göre sıralı)",
            ["Müşteri kodu", "Ayrılma olasılığı", "Risk işaretleri", "Önerilen teklif", "Maliyet", "Beklenen net fayda"],
            [(s["customerID"], yuzde(s["churn_olasiligi"]), s["risk_isaretleri"], s["offer_type"], para(s["offer_cost"]),
              para(s["beklenen_net_fayda"])) for _, s in onerilen.head(10).iterrows()],
            [2.3, 1.7, 6.6, 2.2, 1.3, 1.9], saga=(1, 4, 5))
    r.p(f"Listenin tamamı {A.KAMPANYA_LISTESI.relative_to(A.KOK).as_posix()} dosyasındadır.")


def senaryolar(r, v):
    s, d, egri, ozet = v["senaryolar"], v["duyarlilik"], v["egri"], v["ozet"]
    a, b, c = (s.loc[ad] for ad in (A.SENARYO_A, A.SENARYO_B, A.SENARYO_C))
    r.baslik("9. Senaryo karşılaştırması ve duyarlılık")
    r.baslik("9.1. Senaryolar", 2)
    r.p(f"Üç karar yaklaşımı aynı {tr(len(v['risk']))} müşteri, aynı {para(A.BUTCE)} bütçe ve aynı varsayımlarla karşılaştırılmıştır:")
    r.maddeler(
        "A – Mevcut yöntem: model yoktur. Aylık sözleşmeli ve ilk 12 ayındaki müşterilere, en yüksek ücretliden başlayarak bütçe bitene kadar aynı indirim sunulur.",
        "B – Yalnızca risk modeli: yüksek riskli müşterilere, en riskliden başlayarak bütçe bitene kadar aynı indirim sunulur.",
        "C – Karar destek sistemi: müşteri ve teklif bazında beklenen net fayda, bütçe ve kapasite kısıtlarıyla optimize edilir.",
    )
    r.baslik("9.2. Sonuçlar", 2)
    r.tablo("Senaryo karşılaştırması",
            ["Senaryo", "Teklif sunulan müşteri", "Maliyet", "Beklenen net fayda", "Etiket bazlı net fayda", "Gerçekte ayrılan hedef"],
            [(ad, tr(x["teklif_sunulan_musteri"]), para(x["maliyet"]), para(x["beklenen_net_fayda"]), para(x["etiket_bazli_net_fayda"]),
              f"{tr(x['gercekte_ayrilan_hedef'])} ({yuzde(x['isabet_orani'], 0)})") for ad, x in s.iterrows()],
            [4.3, 2.2, 2.0, 2.5, 2.5, 2.5], saga=(1, 2, 3, 4, 5))
    r.sekil(A.GRAFIK / "08_senaryo_karsilastirma.png", "Senaryolara göre beklenen net fayda")
    r.p(f"Risk modeli tek başına mevcut yönteme göre beklenen net faydayı {para(a['beklenen_net_fayda'])} düzeyinden "
        f"{para(b['beklenen_net_fayda'])} düzeyine çıkarmaktadır; hedeflenen müşterilerden gerçekte ayrılanların payı da "
        f"{yuzde(a['isabet_orani'], 0)} düzeyinden {yuzde(b['isabet_orani'], 0)} düzeyine yükselmektedir. Karar destek sistemi aynı bütçeyle "
        f"{para(c['beklenen_net_fayda'])} beklenen net fayda üretmektedir.")
    ayni_sira = s["beklenen_net_fayda"].rank().equals(s["etiket_bazli_net_fayda"].rank())
    sira_notu = "Senaryoların sıralaması bu ölçümde de değişmemektedir." if ayni_sira else "Bu ölçümde senaryoların sıralaması değişmektedir."
    r.p("Beklenen net fayda modelin kendi olasılıklarıyla hesaplanır. Modelin olasılıkları kararı yanlış yönlendiriyorsa bunu "
        "görebilmek için ikinci bir ölçüm yapılmıştır: aynı etki varsayımı, test müşterilerinin gerçekte ayrılıp ayrılmadığı "
        f"bilgisiyle birleştirilmiştir. {sira_notu} Her iki ölçümde de kampanya etkisi varsayımdır.",
        giris="Etiket bazlı net fayda:")
    destek_payi = ozet["teklif_dagilimi"][A.DESTEK] / ozet["teklif_sunulan_musteri"]
    r.p(f"C senaryosunun üstünlüğünün önemli bir bölümü, A ve B yalnızca indirim sunabilirken sistemin daha ucuz bir teklifi "
        f"seçebilmesinden gelir: önerilen tekliflerin {yuzde(destek_payi, 0)} kadarı teknik destektir ve bu teklifin maliyeti ile "
        "etkisi varsayımdır. Sentetik bir senaryoda sistemin kendi varsayımları altında daha iyi sonuç vermesi, gerçek dünyada "
        "aynı iyileşmenin elde edileceğini kanıtlamaz.", giris="Yorum sınırı:")
    r.baslik("9.3. Bütçe düzeyine duyarlılık", 2)
    r.sekil(A.GRAFIK / "09_butce_egrisi.png", "Bütçe arttıkça senaryoların beklenen net faydası", 14)
    r.p(f"Bütçe arttıkça her ek harcamanın getirisi azalmaktadır. Mevcut yöntemde beklenen net fayda {para(egri[A.SENARYO_A].idxmax())} "
        f"bütçeden sonra düşmektedir: kural, beklenen faydası maliyetini karşılamayan müşterilere de indirim vermeyi sürdürür. "
        f"Karar destek sistemi ise net faydası pozitif olmayan teklifi sunmadığı için bütçenin tamamını harcamak zorunda değildir; "
        f"{para(egri.index.max())} bütçede beklenen net fayda {para(egri[A.SENARYO_C].iloc[-1])} düzeyinde kalır.")
    r.baslik("9.4. Kampanya etkisine duyarlılık", 2)
    r.p("Tekliflerin gerçek etkisi bilinmediği için planlar sabit tutulmuş ve gerçek etkinin varsayılandan düşük ya da yüksek "
        "çıkması durumunda net faydanın ne olacağı hesaplanmıştır.")
    r.tablo(f"Gerçekleşen etki varsayılandan saparsa net fayda ({A.PARA})", ["Senaryo"] + [f"×{tr(x, 2)}" for x in ETKI_DUZEYLERI],
            [(ad,) + tuple(tr(d.loc[x, ad]) for x in ETKI_DUZEYLERI) for ad in d.columns],
            [4.5, 2.3, 2.3, 2.3, 2.3, 2.3], saga=(1, 2, 3, 4, 5))
    r.sekil(A.GRAFIK / "10_duyarlilik.png", "Kampanya etkisi varsayımdan saparsa senaryoların net faydası", 14)
    r.p(f"Gerçek etki varsayılanın yarısı kadar olursa A ve B senaryoları zarar etmekte ({para(d.loc[0.5, A.SENARYO_A])} ve "
        f"{para(d.loc[0.5, A.SENARYO_B])}), önerilen plan ise {para(d.loc[0.5, A.SENARYO_C])} ile başabaşın biraz üzerinde kalmaktadır. "
        f"Önerilen plan, gerçek etki varsayılan etkinin {yuzde(c['basabas_etki_carpani'], 0)} düzeyine inene kadar maliyetini karşılar; "
        f"bu eşik A için {yuzde(a['basabas_etki_carpani'], 0)}, B için {yuzde(b['basabas_etki_carpani'], 0)} düzeyindedir.")


def ekran(r):
    r.baslik("10. Yönetici karar ekranı")
    r.p("Yönetici ekranı Streamlit ile hazırlanmıştır ve modelin teknik ayrıntılarından çok karar için gereken bilgiyi öne "
        "çıkarır. Yönetici sol paneldeki bütçe ve teknik destek kapasitesi değerlerini değiştirdiğinde optimizasyon yeniden "
        "çalışır ve ekrandaki bütün sayılar güncellenir.")
    r.tablo("Karar ekranının bölümleri", ["Bölüm", "Gösterilen bilgi"], [
        ("Öneri", "Önerilen müşteri sayısı ve harcama tek cümleyle; yüksek riskli müşteri sayısı, bütçe kullanımı, beklenen net fayda ve mevcut yönteme göre fark."),
        ("Bütçe nereye gidiyor?", "Teklif türüne göre harcama, müşteri sayısı ve her teklifin kime sunulduğu."),
        ("Bugünkü yöntemle karşılaştırma", "Mevcut yöntem, yalnızca risk sıralaması ve önerilen planın aynı bütçeyle beklenen net faydası."),
        ("Ne kadar emin olabiliriz?", "Teklifler beklenenden az ya da çok işe yararsa planın net faydası; planın zarara geçtiği eşik ve uyarılar."),
        ("Kim risk altında?", "Risk gruplarının dağılımı ve yüksek riskli müşterilerde en sık görülen risk işaretleri."),
        ("Kararınız", "Onayla, değiştir ya da reddet; karar gerekçesiyle ve o andaki varsayımlarla birlikte kayıt dosyasına eklenir."),
        ("Müşteri listesi sekmesi", "Müşteri kodu, ayrılma olasılığı, risk işaretleri, önerilen teklif, maliyet, beklenen net fayda ve belirsizlik; tek bir müşteri için dört seçeneğin karşılaştırması."),
        ("Yöntem ve model sekmesi", "Model ölçütleri, kalibrasyon, grup bazlı kontrol ve senaryoların teknik karşılaştırması."),
    ], [4.6, 11.4])
    if A.EKRAN_GORUNTUSU.exists():
        r.sekil(A.EKRAN_GORUNTUSU, "Yönetici karar ekranının özet sekmesi", 16)
    r.p("Sistem nihai kararı otomatik vermez. Yöneticinin kararı ve gerekçesi "
        f"{A.KARAR_KAYDI.relative_to(A.KOK).as_posix()} dosyasına, kararın verildiği bütçe ve varsayımlarla birlikte kaydedilir.")


def sinirliliklar(r):
    r.baslik("11. Sınırlılıklar, etik ve riskler")
    r.maddeler(
        "Kampanya etkisi doğrulanmamıştır. Teklif maliyetleri, etki oranları, katkı marjı, bütçe ve kapasite sentetik varsayımdır; net fayda sayıları gerçek bir kampanyanın sonucunu göstermez.",
        "Ayrılma riski, müdahalenin etkisi değildir. Yüksek riskli müşteriye teklif sunmak müşteri kaybını kendiliğinden önlemez; etkinin ölçülmesi için kontrol gruplu bir deney gerekir.",
        "Açık eğitim verisi gerçek bir işletmenin güncel müşterilerini temsil etmeyebilir; model başka bir işletmeye doğrudan taşınmamalıdır.",
        "Veri kesitseldir. Tahminler belirli bir gelecek dönem için değil, geçmişte ayrılan müşterilere benzerlik içindir; “gelecek ay ayrılacak müşteri” olarak sunulmamalıdır.",
        "Kampanya kararı test verisindeki müşteriler üzerinde gösterilmiştir. Gerçek kullanımda model, sonucu henüz bilinmeyen güncel müşterilere uygulanır.",
        "Demografik alanlar girdi değildir, ancak diğer girdiler üzerinden dolaylı farklar oluşabilir (Bölüm 6.7). Sonuçlar gruplar bazında izlenmelidir.",
        "Gerçek kişisel veriler kullanılacaksa hukuki dayanak, erişim yetkisi, veri minimizasyonu, saklama süresi ve güvenlik ayrıca ele alınmalıdır.",
        "Sistem öneri sunar; karar yöneticinindir. Öneri gerekçesi, maliyeti ve belirsizliğiyle birlikte gösterilir.",
    )


def sonuc(r, v):
    c = v["senaryolar"].loc[A.SENARYO_C]
    r.baslik("12. Sonuç ve sonraki adımlar")
    r.p("Proje, müşteri kaybı tahmininin tek başına bir karar olmadığını; tahminin teklif maliyeti, beklenen ek etki ve bütçe "
        "kısıtıyla birleştirildiğinde karar desteğine dönüştüğünü göstermektedir. Kurulan sistem görülmemiş veride makul bir "
        "sıralama becerisine sahiptir ve aynı bütçeyle, kendi varsayımları altında, mevcut yönteme göre daha yüksek beklenen "
        f"net fayda ({para(c['beklenen_net_fayda'])}) üretmektedir. Bu sonucun gerçek dünyadaki karşılığı kampanya etkisi ölçülmeden bilinemez.")
    r.p("Önerilen sonraki adımlar:")
    r.maddeler(
        "Planı sınırlı bir pilotla ve rastgele seçilmiş bir kontrol grubuyla uygulamak; teklif türü bazında gerçek etkiyi ölçmek.",
        "Ölçülen etki ve gerçek maliyet verileriyle senaryo tablosunu güncellemek; sentetik değerleri doğrulanmış değerlerle değiştirmek.",
        "Yeterli kampanya verisi oluştuğunda müşteri bazında ek etkiyi doğrudan tahmin eden (uplift) bir model denemek.",
        "Zaman bilgisi içeren veriyle modeli ileriye dönük doğrulamak ve olasılıkların kalibrasyonunu dönemsel olarak izlemek.",
        "Yönetici kararlarının kaydını, önerilerin ne sıklıkla ve hangi gerekçeyle değiştirildiğini görmek için incelemek.",
    )


def kabul_kriterleri(r):
    r.baslik("13. Kabul kriterlerinin karşılanması")
    r.tablo("Yol haritasındaki kabul kriterleri ve karşılandıkları yer", ["Kriter", "Karşılandığı yer"], [
        ("Veri kaynağı ve veri setinin sınırlılıkları açıklanmış olmalı.", "Bölüm 3 ve 11."),
        ("Veri temizleme adımları tekrarlanabilir biçimde kaydedilmiş olmalı.", "Bölüm 4; src/veri_hazirlama.py; outputs/veri_kalite_raporu.md."),
        ("Model, görülmemiş test verisi üzerinde değerlendirilmiş olmalı.", "Bölüm 6.2 ve 6.3."),
        ("Kesinlik, duyarlılık ve en az bir sıralama ölçütü raporlanmalı; mümkünse kalibrasyon incelenmeli.", f"Bölüm 6.3 (Tablo {r.no['olcut']}) ve 6.4."),
        ("Müşteri kaybı tahmini ile kampanya etkisi birbirinden ayrılmış olmalı.", "Bölüm 7.2: etki oranı ayrı bir varsayımdır."),
        ("Teklif maliyeti ve bütçe kısıtı hesaplamaya dahil edilmeli.", "Bölüm 7 ve 8."),
        ("Yönetici ekranında önerinin gerekçesi, maliyeti ve belirsizliği görünmeli.", "Bölüm 10."),
        ("Mevcut yaklaşım ile önerilen yaklaşım aynı varsayımlar altında karşılaştırılmalı.", "Bölüm 9."),
        ("Sentetik veriler ve varsayımlar gerçek işletme sonuçlarından açıkça ayrılmalı.", "Bölüm 7.3; sentetik tablonun dosya adı; ekrandaki uyarılar."),
        ("Kaynakça, kullanılan yazılım ve veri sürümleri rapora eklenmeli.", f"Bölüm 14 ve Tablo {r.no['kaynak']}."),
    ], [9.0, 7.0])


def kaynakca(r):
    r.baslik("14. Kaynakça, yazılım ve veri sürümleri")
    r.baslik("14.1. Kaynakça", 2)
    r.maddeler(
        "Kaggle. Telco Customer Churn veri seti (blastchar). https://www.kaggle.com/datasets/blastchar/telco-customer-churn",
        "IBM. Telco Customer Churn örnek deposu. https://github.com/IBM/telco-customer-churn-on-icp4d",
        "IBM Cognos Analytics. Telco Customer Churn sample. https://www.ibm.com/docs/en/cognos-analytics/12.1.x?topic=samples-telco-customer-churn",
        "scikit-learn. Model evaluation. https://scikit-learn.org/stable/modules/model_evaluation.html",
        "scikit-learn. Metrics API reference. https://scikit-learn.org/stable/api/sklearn.metrics.html",
        "PuLP. Python doğrusal programlama kütüphanesi. https://coin-or.github.io/pulp/",
        "HiGHS. Doğrusal ve karma tam sayılı optimizasyon çözücüsü. https://highs.dev",
        "Streamlit. https://streamlit.io",
    )
    r.baslik("14.2. Yazılım sürümleri", 2)
    r.tablo("Kullanılan yazılım sürümleri", ["Yazılım", "Sürüm"],
            [("Python", sys.version.split()[0])] + [(ad, version(ad)) for ad in YAZILIMLAR], [6.0, 4.0])
    r.p(f"Veri sürümü, kaynak kaydındaki (Tablo {r.no['kaynak']}) dosya tarihi ve SHA-256 özetiyle kaydedilmiştir.")
    r.baslik("14.3. Çalıştırma ve dosya düzeni", 2)
    r.p("Bütün çıktılar ham veriden aşağıdaki komutlarla yeniden üretilebilir:")
    for komut in ("py -m pip install -r requirements.txt", "py src/hepsini_calistir.py", "py -m streamlit run app.py"):
        r.p(komut, girinti=True).paragraph_format.space_after = Pt(2)
    r.tablo("Proje klasör yapısı", ["Klasör / dosya", "İçerik"], [
        ("data/raw/", "Kaynaktan indirilen, değiştirilmemiş ham CSV."),
        ("data/processed/", "Temizlenmiş veri ve sentetik teklif senaryosu tablosu."),
        ("src/", "Veri hazırlama, keşifsel analiz, model, teklif senaryosu, optimizasyon, senaryo karşılaştırması ve rapor betikleri; varsayımlar ayarlar.py dosyasında."),
        ("app.py", "Yönetici karar ekranı."),
        ("outputs/", "Grafikler, tablolar, model ölçütleri, müşteri risk tablosu, kampanya listesi, yönetici özeti ve karar kaydı."),
        ("docs/", "Veri sözlüğü ve bu rapor."),
    ], [4.2, 11.8])


def main():
    baslik_yaz("Proje raporu")
    v = verileri_topla()
    r = Rapor()
    kapak_ve_icindekiler(r)
    yonetici_ozeti(r, v)
    problem(r)
    veri(r, v)
    kalite(r, v)
    kesif(r, v)
    model(r, v)
    teklif(r, v)
    optimizasyon(r, v)
    senaryolar(r, v)
    ekran(r)
    sinirliliklar(r)
    sonuc(r, v)
    kabul_kriterleri(r)
    kaynakca(r)

    A.RAPOR.parent.mkdir(parents=True, exist_ok=True)
    try:
        r.doc.save(str(A.RAPOR))
    except PermissionError:
        sys.exit(f"Rapor yazılamadı: {A.RAPOR.name} Word'de açık olabilir. Dosyayı kapatıp yeniden deneyin.")
    print(f"Yazıldı: {A.RAPOR.relative_to(A.KOK)}  ({r.sekil_no} şekil, {r.tablo_no} tablo)")
    if word_ile_tamamla(A.RAPOR):
        print(f"İçindekiler güncellendi; PDF kopyası: {A.RAPOR.with_suffix('.pdf').relative_to(A.KOK)}")
    else:
        print("Word bulunamadı ya da açılamadı: içindekiler tablosunu Word'de sağ tıklayıp güncelleyin.")


if __name__ == "__main__":
    main()
